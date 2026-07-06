from flask import Blueprint, render_template, redirect, url_for, flash, Response
from routes.auth import login_required, role_required
from database.connection import execute_query
import openpyxl
from io import BytesIO

reports_bp = Blueprint('reports', __name__)

@reports_bp.route('/admin/reports')
@login_required
@role_required(['Admin', 'SuperAdmin'])
def index():
    # Renders the reports list and summaries
    try:
        # Total counts
        total_books = execute_query("SELECT COUNT(*) as cnt FROM Books")[0]['cnt']
        borrowed_books = execute_query("SELECT COUNT(*) as cnt FROM BorrowedBooks WHERE Status = 'Approved'")[0]['cnt']
        returned_books = execute_query("SELECT COUNT(*) as cnt FROM BorrowedBooks WHERE Status = 'Returned'")[0]['cnt']
        
        # Product Sales
        sales_data = execute_query(
            "SELECT SUM(TotalAmount) as revenue, COUNT(*) as cnt FROM Orders WHERE OrderStatus = 'Delivered'"
        )[0]
        total_revenue = sales_data['revenue'] or 0.00
        completed_sales = sales_data['cnt']
        
        # Support ticket stats
        open_tickets = execute_query("SELECT COUNT(*) as cnt FROM SupportTickets WHERE Status = 'Open'")[0]['cnt']
        resolved_tickets = execute_query("SELECT COUNT(*) as cnt FROM SupportTickets WHERE Status = 'Resolved'")[0]['cnt']
        
        # Active users count
        active_users = execute_query("SELECT COUNT(*) as cnt FROM Users WHERE UserStatus = 'Active'")[0]['cnt']
        
        return render_template(
            'reports.html',
            total_books=total_books,
            borrowed_books=borrowed_books,
            returned_books=returned_books,
            total_revenue=total_revenue,
            completed_sales=completed_sales,
            open_tickets=open_tickets,
            resolved_tickets=resolved_tickets,
            active_users=active_users
        )
    except Exception as e:
        flash(f"Error loading reports summary: {e}", "danger")
        return redirect(url_for('dashboard.index'))

@reports_bp.route('/admin/reports/excel')
@login_required
@role_required(['Admin', 'SuperAdmin'])
def export_excel():
    # Generates a dynamic multi-sheet Excel spreadsheet using openpyxl
    wb = openpyxl.Workbook()
    
    try:
        # Sheet 1: Library Inventory
        ws_books = wb.active
        ws_books.title = "Books Inventory"
        books = execute_query(
            """SELECT b.ISBN, b.Title, b.Author, b.Edition, b.Publisher, b.PublicationYear, 
                      c.CategoryName, b.TotalCopies, b.AvailableCopies 
               FROM Books b
               INNER JOIN Categories c ON b.CategoryID = c.CategoryID
               ORDER BY b.Title ASC"""
        )
        ws_books.append(["ISBN", "Title", "Author", "Edition", "Publisher", "Year", "Category", "Total Copies", "Available Copies"])
        for b in books:
            ws_books.append([b['ISBN'], b['Title'], b['Author'], b['Edition'], b['Publisher'], b['PublicationYear'], b['CategoryName'], b['TotalCopies'], b['AvailableCopies']])
            
        # Sheet 2: Borrowing History
        ws_borrows = wb.create_sheet(title="Borrow Log")
        borrows = execute_query(
            """SELECT bb.BorrowID, u.Username, u.CollegeName, b.Title, b.ISBN, bb.BorrowDate, bb.DueDate, bb.ReturnDate, bb.Status 
               FROM BorrowedBooks bb
               INNER JOIN Users u ON bb.UserID = u.UserID
               INNER JOIN Books b ON bb.BookID = b.BookID
               ORDER BY bb.BorrowDate DESC"""
        )
        ws_borrows.append(["Borrow ID", "Student Username", "Campus College", "Book Title", "ISBN", "Borrow Date", "Due Date", "Return Date", "Status"])
        for br in borrows:
            ws_borrows.append([br['BorrowID'], br['Username'], br['CollegeName'], br['Title'], br['ISBN'], br['BorrowDate'], br['DueDate'], br['ReturnDate'], br['Status']])
            
        # Sheet 3: Sales Orders
        ws_orders = wb.create_sheet(title="Store Orders")
        orders = execute_query(
            """SELECT o.OrderID, u.Username, u.CollegeName, o.OrderDate, o.TotalAmount, o.ShippingAddress, o.PaymentStatus, o.OrderStatus 
               FROM Orders o
               INNER JOIN Users u ON o.BuyerID = u.UserID
               ORDER BY o.OrderDate DESC"""
        )
        ws_orders.append(["Order ID", "Buyer Username", "College", "Order Date", "Amount", "Shipping Address", "Payment Status", "Order Status"])
        for o in orders:
            ws_orders.append([o['OrderID'], o['Username'], o['CollegeName'], o['OrderDate'], o['TotalAmount'], o['ShippingAddress'], o['PaymentStatus'], o['OrderStatus']])
            
        # Save workbook to memory buffer
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return Response(
            output.read(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-disposition": "attachment; filename=GyanPustak_System_Report.xlsx"}
        )
    except Exception as e:
        flash(f"Excel generation failed: {e}", "danger")
        return redirect(url_for('reports.index'))
